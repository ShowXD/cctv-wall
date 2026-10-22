"""攝影機清單與試算表（.xlsx／.csv）互轉。

現場拿到的連線資訊欄位名稱與順序每份都不同，所以匯入分兩步：先讀出原始儲存格，
再依使用者指認的「第幾欄是什麼」轉成攝影機。轉換本身交給 Camera.from_raw，
網址優先、否則 IP＋Port＋廠牌樣板、網址內嵌帳密等規則都和設定檔一致。
"""
from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from cctv_wall.models import NAME_LIMIT, Camera, Site, Workspace

FIELDS = ("site", "name", "url", "host", "port", "stream", "user", "password")
EXPORT_FIELDS = ("site", "name", "url", "user", "password")
CSV_ENCODINGS = ("utf-8-sig", "cp950")   # cp950：台灣 Excel「另存 CSV」的預設編碼

# 依序比對：前面的欄位先認領，"username" 才不會被當成名稱、"串流網址" 不會被當成串流編號。
_HEADER_WORDS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("site", ("地點", "場所", "站點", "據點", "site", "location")),
    ("password", ("密碼", "password", "pwd", "pass")),
    ("user", ("帳號", "使用者", "用戶", "user", "account", "login")),
    ("url", ("網址", "url", "rtsp")),
    ("host", ("ip", "host", "主機", "位址")),
    ("port", ("port", "埠")),
    ("stream", ("串流", "碼流", "stream")),
    ("name", ("名稱", "名字", "name", "攝影機", "camera", "鏡頭")),
)


def _cell(value: Any) -> str:
    """Excel 的數字一律是浮點數，554.0 要寫回 554 才是合法的 Port。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _open_workbook(path: Path):
    from openpyxl import load_workbook  # 只有匯入 Excel 時才載入
    from openpyxl.utils.exceptions import InvalidFileException
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError) as error:
        # 副檔名是 .xlsx 但內容不是（或已損毀）：當成讀檔錯誤回報，不讓整個指令崩潰。
        raise ValueError(f"{path.name}: not a valid .xlsx workbook") from error


def sheet_names(path: str | Path) -> tuple[list[str], int]:
    """活頁簿的工作表名稱與作用中工作表的位置；CSV 沒有工作表，回傳 ([], 0)。"""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        return [], 0
    workbook = _open_workbook(path)
    try:
        return list(workbook.sheetnames), workbook.sheetnames.index(workbook.active.title)
    finally:
        workbook.close()


def read_rows(path: str | Path, sheet: str | None = None) -> list[list[str]]:
    """讀出指定工作表（預設作用中的那張）或 CSV 的所有列，儲存格一律是去掉前後空白的字串。"""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        workbook = _open_workbook(path)
        try:
            worksheet = workbook[sheet] if sheet else workbook.active
            return [[_cell(value) for value in row] for row in worksheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    raw = path.read_bytes()
    for encoding in CSV_ENCODINGS:
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"{path.name}: not UTF-8 or Big5 (cp950) text")
    return [[_cell(value) for value in row] for row in csv.reader(io.StringIO(text))]


def guess_mapping(header: Sequence[str]) -> dict[str, int]:
    """依標題文字猜「欄位 → 第幾欄」，猜不到的欄位不放；只是初值，使用者會再確認。"""
    mapping: dict[str, int] = {}
    for column, title in enumerate(header):
        text = title.lower().replace(" ", "")
        for field, words in _HEADER_WORDS:
            if field not in mapping and any(word in text for word in words):
                mapping[field] = column
                break
    return mapping


def cameras_from_rows(rows: Sequence[Sequence[str]], mapping: Mapping[str, int], vendor: str,
                      default_site: str, first_row: int = 1
                      ) -> tuple[list[tuple[str, Camera]], list[tuple[int, str]]]:
    """把資料列轉成（地點, 攝影機）。全空白的列略過；有錯的列記下（列號, 原因）後繼續，
    整份檔不會因為一列打錯而匯不進來。`first_row` 是 rows[0] 在試算表上的列號。"""
    cameras: list[tuple[str, Camera]] = []
    errors: list[tuple[int, str]] = []
    for offset, row in enumerate(rows):
        raw = {field: row[column] if column < len(row) else ""
               for field, column in mapping.items()}
        if not any(raw.values()):
            continue
        site = (raw.pop("site", "") or default_site).strip()[:NAME_LIMIT]
        try:
            if not site:
                raise ValueError("site needs a name")
            cameras.append((site, Camera.from_raw({**raw, "vendor": vendor})))
        except ValueError as error:
            errors.append((first_row + offset, str(error)))
    return cameras, errors


def merge_into(workspace: Workspace, imported: Iterable[tuple[str, Camera]]) -> tuple[Workspace, int]:
    """附加到同名地點（沒有就新建），回傳新的工作區與實際加入的台數。

    同一地點已有相同網址＋帳號的攝影機就略過：同一份清單匯入兩次不會變成兩倍。
    有 AI 連線的地點拒絕重名攝影機時，Site 的驗證會拋 ValueError。
    """
    updated = workspace
    added = 0
    for site_name, camera in imported:
        site = updated.site(site_name)
        if site is None:
            updated = updated.with_sites((*updated.sites, Site(site_name, (camera,))))
        elif any(c.url == camera.url and c.user == camera.user for c in site.cameras):
            continue
        else:
            updated = updated.with_site(site_name, site.with_cameras((*site.cameras, camera)))
        added += 1
    return updated, added


def table_rows(sites: Iterable[Site], include_credentials: bool) -> list[list[str]]:
    """匯出用的資料列，欄位順序同 EXPORT_FIELDS；不含帳密時只到網址。"""
    fields = EXPORT_FIELDS if include_credentials else EXPORT_FIELDS[:3]
    return [[site.name if field == "site" else getattr(camera, field) for field in fields]
            for site in sites for camera in site.cameras]


def write_rows(path: str | Path, rows: Iterable[Sequence[str]]) -> None:
    """.xlsx 寫成單一工作表；其他副檔名寫 UTF-8（含 BOM）CSV，Excel 直接開中文不亂碼。"""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        workbook = Workbook()
        for row in rows:
            workbook.active.append(list(row))
        workbook.save(path)
        return
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        csv.writer(file).writerows(rows)
